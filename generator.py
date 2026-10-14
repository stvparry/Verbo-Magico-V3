from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT = "D9EAF7"
WHITE = "FFFFFF"


def normalise(payload):

    if "value" not in payload:
        raise ValueError(
            f"Expected 'value' key. Found: {list(payload.keys())}"
        )

    value = payload["value"]

    if "moods" not in value:
        raise ValueError(
            f"Expected 'moods' key. Found: {list(value.keys())}"
        )

    return value

def strip_pronoun(text):

    pronouns = [
        "yo",
        "tú",
        "vos",
        "él",
        "ella",
        "usted",
        "nosotros",
        "nosotras",
        "vosotros",
        "vosotras",
        "ellos",
        "ellas",
        "ustedes",
    ]

    text = text.strip()

    for p in sorted(pronouns, key=len, reverse=True):

        prefix = p + " "

        if text.lower().startswith(prefix):
            return text[len(prefix):]

    return text


def get_forms(data, mood, tense):

    entries = data["moods"][mood][tense]

    result = {}

    for entry in entries:

        pronoun = entry.get("pr")

        if not pronoun:
            continue

        result[pronoun] = strip_pronoun(
            entry["c"][0]
        )

    return result

    entries = data["moods"][mood][tense]

    result = {}

    for entry in entries:

        pronoun = entry.get("pr")

        if not pronoun:
            continue

        result[pronoun] = strip_pronoun(
            entry["c"][0]
        )

    return result


def style_title(ws, text):

    ws.merge_cells("A1:G1")

    c = ws["A1"]

    c.value = text

    c.font = Font(
        size=18,
        bold=True,
        color=WHITE,
    )

    c.fill = PatternFill(
        "solid",
        fgColor=NAVY,
    )

    c.alignment = Alignment(
        horizontal="center",
    )


def build_workbook(verb, payload):

    data = normalise(payload)

    wb = Workbook()

    ws = wb.active
    ws.title = "Quick reference"

    style_title(
        ws,
        f"{verb.upper()} conjugation",
    )

    ws["A3"] = "Infinitive"
    ws["B3"] = verb

    headers = [
        "Person",
        "Present",
    ]

    for col, header in enumerate(headers, start=1):

        c = ws.cell(6, col)

        c.value = header

        c.fill = PatternFill(
            "solid",
            fgColor=LIGHT,
        )

        c.font = Font(
            bold=True
        )

    present = get_forms(
        data,
        "indicativo",
        "presente",
    )

    row = 7

    persons = [
        "yo",
        "tú",
        "él",
        "nosotros",
        "vosotros",
        "ellos",
    ]

    for pronoun in persons:

        ws.cell(row, 1).value = pronoun

        ws.cell(row, 2).value = present.get(
            pronoun,
            ""
        )

        row += 1

    output = BytesIO()

    wb.save(output)

    output.seek(0)

    return output.getvalue()
